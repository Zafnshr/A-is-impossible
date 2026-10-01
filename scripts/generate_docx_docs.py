import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_styled_document(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Page Margins: 1 inch everywhere
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B) # Slate-800

    in_code_block = False
    code_lines = []

    in_table = False
    table_lines = []

    def flush_code_block(c_lines):
        if not c_lines:
            return
        code_text = "".join(c_lines).rstrip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)

        # Style background using XML shading
        pPr = p._p.get_or_add_pPr()
        shd = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
        pPr.append(shd)

        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    def flush_table(t_lines):
        if not t_lines:
            return
        rows_data = []
        for l in t_lines:
            l = l.strip()
            if not l.startswith('|'):
                continue
            cells = [c.strip() for c in l.split('|')[1:-1]]
            if all(re.match(r'^[\:\-\s]+$', c) for c in cells):
                continue # delimiter row
            rows_data.append(cells)

        if not rows_data:
            return

        col_count = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=col_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for r_idx, row in enumerate(rows_data):
            for c_idx, cell_value in enumerate(row):
                if c_idx < col_count:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_value
                    # Formatting
                    tcPr = cell._tc.get_or_add_tcPr()
                    if r_idx == 0:
                        # Header background
                        shd = parse_xml(r'<w:shd {} w:fill="E2E8F0"/>'.format(nsdecls('w')))
                        tcPr.append(shd)
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.bold = True
                                r.font.name = 'Calibri'
                                r.font.size = Pt(10)
                                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                    else:
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.name = 'Calibri'
                                r.font.size = Pt(9.5)

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_after = Pt(6)

    for line in lines:
        stripped = line.strip()

        # Handle Code Blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                flush_code_block(code_lines)
                code_lines = []
            else:
                if in_table:
                    in_table = False
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle Tables
        if stripped.startswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        else:
            if in_table:
                in_table = False
                flush_table(table_lines)
                table_lines = []

        # Horizontal rule
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = parse_xml(r'<w:pBdr {}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>'.format(nsdecls('w')))
            pPr.append(pbdr)
            p.paragraph_format.space_after = Pt(8)
            continue

        # Headings
        if stripped.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(stripped[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x08, 0x91, 0xB2) # Cyan-600
            continue
        elif stripped.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(stripped[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0E, 0x74, 0x90) # Cyan-700
            continue
        elif stripped.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(stripped[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            continue
        elif stripped.startswith('#### '):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            run = h.add_run(stripped[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            continue

        # Callouts (> [!IMPORTANT] ...)
        if stripped.startswith('> '):
            quote_text = stripped[2:]
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pbdr = parse_xml(r'<w:pBdr {}><w:left w:val="single" w:sz="24" w:space="8" w:color="0891B2"/></w:pBdr>'.format(nsdecls('w')))
            pPr.append(pbdr)
            r = p.add_run(quote_text)
            r.font.italic = True
            r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            continue

        # Bullet lists
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            text = stripped[2:]
            # Format bold (**...**)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
            continue

        # Numbered lists
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            text = num_match.group(2)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
            continue

        # Regular paragraphs
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            parts = re.split(r'(\*\*.*?\*\*)', line.rstrip())
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)

    doc.save(docx_path)
    print(f"Successfully generated DOCX at {docx_path}")

if __name__ == '__main__':
    md_file = r'c:\Users\abdal\Downloads\A is impossible\IMPORT_DOCUMENTATION.md'
    docx_file = r'c:\Users\abdal\Downloads\A is impossible\IMPORT_DOCUMENTATION.docx'
    create_styled_document(md_file, docx_file)
